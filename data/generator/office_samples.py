"""Documentos sintéticos en formatos de oficina (Word, Excel, PowerPoint, CSV, TXT, imagen)
y versiones vacías de cada formato para probar cómo el sistema maneja archivos sin contenido.
"""

import csv
import io
import random
from dataclasses import asdict, dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import docx
import openpyxl
import pptx
import pymupdf
from openpyxl.styles import Font
from pptx.util import Inches, Pt

CATEGORIES = ["Viajes", "Software", "Oficina", "Marketing", "Capacitación"]
AREAS = ["Finanzas", "Operaciones", "Ventas", "Tecnología"]


@dataclass
class Expense:
    date: str
    area: str
    category: str
    description: str
    amount: float


def _expenses(rng: random.Random, today: date, count: int = 24) -> list[Expense]:
    items = []
    for _ in range(count):
        day = today - timedelta(days=rng.randint(1, 90))
        category = rng.choice(CATEGORIES)
        items.append(
            Expense(
                date=day.isoformat(),
                area=rng.choice(AREAS),
                category=category,
                description=f"Gasto de {category.lower()} #{rng.randint(100, 999)}",
                amount=round(rng.uniform(80, 2500), 2),
            )
        )
    return sorted(items, key=lambda e: e.date)


def make_policy_docx(path: Path) -> dict[str, Any]:
    document = docx.Document()
    document.add_heading("Política de rendición de gastos 2026", level=1)
    document.add_paragraph(
        "Esta política aplica a todos los colaboradores de Comercial Demo SpA que incurran en "
        "gastos por cuenta de la empresa. Documento sintético generado para pruebas."
    )
    document.add_heading("1. Plazos", level=2)
    document.add_paragraph(
        "Los gastos deben rendirse dentro de los 15 días corridos siguientes a la fecha del gasto."
    )
    document.add_heading("2. Límites por categoría", level=2)
    limits = [
        ("Viajes", "1.500"),
        ("Software", "800"),
        ("Capacitación", "2.000"),
        ("Oficina", "300"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Categoría"
    table.rows[0].cells[1].text = "Límite por gasto (USD)"
    for category, limit in limits:
        row = table.add_row()
        row.cells[0].text = category
        row.cells[1].text = limit
    document.add_heading("3. Aprobaciones", level=2)
    document.add_paragraph(
        "Todo gasto que supere el límite de su categoría requiere aprobación previa del gerente "
        "de área y del área de Finanzas."
    )
    document.save(path)
    return {"filename": path.name, "rendition_days": 15, "limits": dict(limits)}


def make_expenses_xlsx(path: Path, expenses: list[Expense]) -> dict[str, Any]:
    workbook = openpyxl.Workbook()
    detail = workbook.active
    assert detail is not None
    detail.title = "Gastos"
    detail.append(["Fecha", "Área", "Categoría", "Descripción", "Monto (USD)"])
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for e in expenses:
        detail.append([e.date, e.area, e.category, e.description, e.amount])

    summary = workbook.create_sheet("Presupuesto")
    summary.append(["Área", "Presupuesto trimestral (USD)", "Gastado (USD)"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    budget = {
        "Finanzas": 8000,
        "Operaciones": 12000,
        "Ventas": 10000,
        "Tecnología": 15000,
    }
    spent = {
        area: round(sum(e.amount for e in expenses if e.area == area), 2)
        for area in AREAS
    }
    for area in AREAS:
        summary.append([area, budget[area], spent[area]])
    workbook.save(path)
    return {
        "filename": path.name,
        "sheets": ["Gastos", "Presupuesto"],
        "budget": budget,
        "spent": spent,
    }


def make_results_pptx(path: Path, rng: random.Random) -> dict[str, Any]:
    presentation = pptx.Presentation()
    revenue = {q: round(rng.uniform(180_000, 320_000), 2) for q in ["Q1", "Q2", "Q3"]}

    title = presentation.slides.add_slide(presentation.slide_layouts[0])
    title.shapes.title.text = "Resultados trimestrales 2026"
    title.placeholders[1].text = "Comercial Demo SpA — Documento sintético para pruebas"

    for quarter, amount in revenue.items():
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Ingresos {quarter}"
        body = slide.placeholders[1].text_frame
        body.text = f"Ingresos totales: {amount:,.2f} USD"
        body.add_paragraph().text = f"Margen bruto estimado: {rng.randint(28, 41)}%"
        slide.notes_slide.notes_text_frame.text = f"Comentar la variación de {quarter}."

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "Resumen por trimestre"
    shape = table_slide.shapes.add_table(
        4, 2, Inches(1), Inches(2), Inches(6), Inches(2)
    )
    shape.table.cell(0, 0).text = "Trimestre"
    shape.table.cell(0, 1).text = "Ingresos (USD)"
    for row, (quarter, amount) in enumerate(revenue.items(), start=1):
        shape.table.cell(row, 0).text = quarter
        shape.table.cell(row, 1).text = f"{amount:,.2f}"
        for col in range(2):
            for paragraph in shape.table.cell(row, col).text_frame.paragraphs:
                paragraph.font.size = Pt(14)
    presentation.save(path)
    return {
        "filename": path.name,
        "slides": len(presentation.slides),
        "revenue": revenue,
    }


def make_movements_csv(path: Path, expenses: list[Expense]) -> dict[str, Any]:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["fecha", "area", "categoria", "descripcion", "monto_usd"])
        for e in expenses:
            writer.writerow(
                [e.date, e.area, e.category, e.description, f"{e.amount:.2f}"]
            )
    return {"filename": path.name, "rows": len(expenses), "delimiter": ";"}


def make_meeting_txt(path: Path, today: date) -> dict[str, Any]:
    next_review = today + timedelta(days=30)
    path.write_text(
        "Notas de reunión — Comité de Finanzas\n"
        f"Fecha: {today.isoformat()}\n\n"
        "Asistentes: Gerencia de Finanzas, Operaciones y Tecnología.\n\n"
        "Acuerdos:\n"
        "1. Reducir el gasto en software un 10% renegociando licencias anuales.\n"
        "2. Toda compra sobre 5.000 USD requerirá orden de compra aprobada.\n"
        f"3. Próxima revisión del presupuesto: {next_review.isoformat()}.\n\n"
        "Documento sintético generado para pruebas.\n",
        encoding="utf-8",
    )
    return {"filename": path.name, "next_review": next_review.isoformat()}


def make_invoice_image(path: Path, source_pdf: Path) -> dict[str, Any]:
    """Una factura como foto/escaneo JPG (sin texto extraíble: requiere OCR)."""
    with pymupdf.open(source_pdf) as pdf:
        pix = pdf[0].get_pixmap(dpi=110)
        path.write_bytes(pix.tobytes("jpg", jpg_quality=80))
    return {"filename": path.name, "source": source_pdf.name}


def make_empty_files(folder: Path) -> list[str]:
    """Un archivo sin contenido por formato: el sistema debe marcarlos como fallidos con un
    mensaje claro ("No se encontró texto"), sin romperse."""
    folder.mkdir(parents=True, exist_ok=True)
    docx.Document().save(folder / "vacio.docx")
    openpyxl.Workbook().save(folder / "vacio.xlsx")
    pptx.Presentation().save(folder / "vacio.pptx")
    (folder / "vacio.csv").write_text("", encoding="utf-8")
    (folder / "vacio.txt").write_text("", encoding="utf-8")
    blank = pymupdf.open()
    blank.new_page()
    buffer = io.BytesIO()
    blank.save(buffer)
    (folder / "vacio.pdf").write_bytes(buffer.getvalue())
    return sorted(p.name for p in folder.iterdir())


def _docx(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    document = docx.Document()
    document.add_heading(title, level=1)
    for heading, body in sections:
        document.add_heading(heading, level=2)
        document.add_paragraph(body)
    document.add_paragraph("Documento sintético generado para pruebas.")
    document.save(path)


def make_payments_procedure_docx(path: Path) -> dict[str, Any]:
    _docx(
        path,
        "Procedimiento de pago a proveedores",
        [
            (
                "1. Recepción",
                "Toda factura se registra en el sistema dentro de 2 días hábiles.",
            ),
            (
                "2. Validación",
                (
                    "Finanzas valida que la factura tenga orden de compra asociada y que el monto "
                    "coincida con lo recibido. Las facturas sobre 5.000 USD sin orden de compra se "
                    "rechazan."
                ),
            ),
            (
                "3. Pago",
                "Los pagos se ejecutan los días martes y jueves por transferencia.",
            ),
        ],
    )
    return {
        "filename": path.name,
        "po_threshold": 5000,
        "payment_days": ["martes", "jueves"],
    }


def make_closing_memo_docx(path: Path, today: date) -> dict[str, Any]:
    deadline = (today.replace(day=1) + timedelta(days=40)).replace(day=5)
    _docx(
        path,
        "Memo: cierre contable mensual",
        [
            (
                "Plazo",
                f"El cierre del mes debe quedar listo a más tardar el {deadline.isoformat()}.",
            ),
            (
                "Responsables",
                (
                    "Contabilidad concilia bancos; Tesorería informa pagos pendientes; cada gerencia "
                    "aprueba sus provisiones."
                ),
            ),
        ],
    )
    return {"filename": path.name, "deadline": deadline.isoformat()}


def make_annual_budget_xlsx(path: Path, rng: random.Random) -> dict[str, Any]:
    months = [
        "Ene",
        "Feb",
        "Mar",
        "Abr",
        "May",
        "Jun",
        "Jul",
        "Ago",
        "Sep",
        "Oct",
        "Nov",
        "Dic",
    ]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Presupuesto 2026"
    sheet.append(["Área", *months, "Total"])
    totals = {}
    for area in AREAS:
        values = [round(rng.uniform(2_000, 6_000), 2) for _ in months]
        totals[area] = round(sum(values), 2)
        sheet.append([area, *values, totals[area]])
    workbook.save(path)
    return {"filename": path.name, "totals": totals}


def make_reconciliation_xlsx(path: Path, rng: random.Random) -> dict[str, Any]:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Conciliación"
    sheet.append(["Concepto", "Monto (USD)"])
    bank = round(rng.uniform(30_000, 50_000), 2)
    in_transit = round(rng.uniform(500, 3_000), 2)
    uncashed = round(rng.uniform(300, 2_000), 2)
    books = round(bank + in_transit - uncashed, 2)
    for row in [
        ("Saldo según banco", bank),
        ("(+) Depósitos en tránsito", in_transit),
        ("(-) Cheques girados no cobrados", uncashed),
        ("Saldo según libros", books),
    ]:
        sheet.append(list(row))
    workbook.save(path)
    return {"filename": path.name, "bank": bank, "books": books}


def make_commercial_plan_pptx(path: Path) -> dict[str, Any]:
    presentation = pptx.Presentation()
    slides = [
        (
            "Plan comercial 2027",
            "Comercial Demo SpA — Documento sintético para pruebas",
        ),
        (
            "Objetivos",
            "Crecer 15% en ventas\nAbrir 2 nuevas regiones\nReducir el churn al 4%",
        ),
        ("Presupuesto", "Marketing: 120.000 USD\nFuerza de ventas: 340.000 USD"),
    ]
    for index, (title, body) in enumerate(slides):
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[0 if index == 0 else 1]
        )
        slide.shapes.title.text = title
        slide.placeholders[1].text_frame.text = body
    presentation.save(path)
    return {"filename": path.name, "growth_target": "15%"}


def make_vendors_csv(path: Path) -> dict[str, Any]:
    from generate import VENDORS  # evita import circular al cargar el módulo

    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["proveedor", "rut", "plazo_pago_dias", "categoria"])
        for i, (name, tax_id) in enumerate(VENDORS):
            writer.writerow(
                [name, tax_id, [15, 30, 45][i % 3], CATEGORIES[i % len(CATEGORIES)]]
            )
    return {"filename": path.name, "rows": len(VENDORS)}


def make_audit_minutes_txt(path: Path, today: date) -> dict[str, Any]:
    path.write_text(
        "Acta — Comité de Auditoría\n"
        f"Fecha: {(today - timedelta(days=12)).isoformat()}\n\n"
        "Hallazgos:\n"
        "1. Tres facturas sin orden de compra superaron los 5.000 USD.\n"
        "2. La conciliación bancaria de agosto tuvo una diferencia de 1.240 USD, ya aclarada.\n\n"
        "Compromisos: reforzar el control de órdenes de compra antes del próximo trimestre.\n\n"
        "Documento sintético generado para pruebas.\n",
        encoding="utf-8",
    )
    return {"filename": path.name, "findings": 2}


def generate_office_samples(
    out: Path, rng: random.Random, today: date, extended: bool = False
) -> dict[str, Any]:
    expenses = _expenses(rng, today)
    truth: dict[str, Any] = {
        "policy_docx": make_policy_docx(out / "politica_gastos.docx"),
        "expenses_xlsx": make_expenses_xlsx(out / "gastos_trimestre.xlsx", expenses),
        "results_pptx": make_results_pptx(out / "resultados_trimestrales.pptx", rng),
        "movements_csv": make_movements_csv(out / "movimientos_gastos.csv", expenses),
        "meeting_txt": make_meeting_txt(out / "notas_reunion_finanzas.txt", today),
        "invoice_image": make_invoice_image(
            out / "factura_foto.jpg", out / "factura_0002.pdf"
        ),
        "expenses": [asdict(e) for e in expenses],
        "empty_files": make_empty_files(out / "vacios"),
    }
    if extended:
        truth["extended"] = {
            "payments_procedure": make_payments_procedure_docx(
                out / "procedimiento_pago_proveedores.docx"
            ),
            "closing_memo": make_closing_memo_docx(
                out / "memo_cierre_mensual.docx", today
            ),
            "annual_budget": make_annual_budget_xlsx(
                out / "presupuesto_anual_2026.xlsx", rng
            ),
            "reconciliation": make_reconciliation_xlsx(
                out / "conciliacion_bancaria.xlsx", rng
            ),
            "commercial_plan": make_commercial_plan_pptx(
                out / "plan_comercial_2027.pptx"
            ),
            "vendors": make_vendors_csv(out / "proveedores.csv"),
            "audit_minutes": make_audit_minutes_txt(
                out / "acta_comite_auditoria.txt", today
            ),
            "invoice_image_2": make_invoice_image(
                out / "factura_foto_2.jpg", out / "factura_0005.pdf"
            ),
        }
    return truth
