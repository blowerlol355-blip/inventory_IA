import io

import docx
import openpyxl
import pptx
import pytest

from app.services.ingestion.filetypes import CSV, DOCX, PDF, PPTX, TXT, XLSX, detect_content_type
from app.services.ingestion.parsing import ParsingError, parse_document
from tests.conftest import make_pdf, make_scanned_pdf


def test_pdf_keeps_page_numbers() -> None:
    data = make_pdf(
        [
            "<p>Factura número F-1 del proveedor Acme</p>",
            "<p>Total a pagar de la factura: 1.190</p>",
        ]
    )
    parsed = parse_document(data, PDF)
    assert parsed.page_count == 2
    assert [p.page for p in parsed.pages] == [1, 2]
    assert "Acme" in parsed.pages[0].text
    assert "1.190" in parsed.pages[1].text
    assert parsed.ocr_pages == []


def test_docx_paragraphs_and_tables() -> None:
    document = docx.Document()
    document.add_paragraph("Contrato de servicios con Proveedor Uno")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Monto"
    table.rows[0].cells[1].text = "5.000"
    buffer = io.BytesIO()
    document.save(buffer)

    data = buffer.getvalue()
    assert detect_content_type(data) == DOCX
    parsed = parse_document(data, DOCX)
    assert "Proveedor Uno" in parsed.full_text
    assert "Monto | 5.000" in parsed.full_text


def test_ligatures_are_normalized() -> None:
    parsed = parse_document(make_pdf(["<p>Cláusula de conﬁdencialidad ﬁrmada</p>"]), PDF)
    assert "confidencialidad firmada" in parsed.full_text


def test_detect_content_type_uses_bytes_not_extension() -> None:
    assert detect_content_type(make_pdf(["<p>hola mundo desde un pdf</p>"])) == PDF
    assert detect_content_type(b"MZ\x90\x00 ejecutable disfrazado de pdf") is None
    assert detect_content_type(b"texto plano") is None


def test_corrupt_pdf_raises_parsing_error() -> None:
    with pytest.raises(ParsingError):
        parse_document(b"%PDF-1.7 esto no es un pdf real", PDF)


def test_scanned_page_is_marked_for_ocr_but_blank_page_is_not() -> None:
    parsed = parse_document(make_scanned_pdf("Factura escaneada número 123"), PDF)
    assert [p.page for p in parsed.ocr_pages] == [1]
    image = parsed.ocr_pages[0].image
    assert image is not None and image[1:4] == b"PNG"

    blank = parse_document(make_pdf([""]), PDF)
    assert blank.ocr_pages == []  # una página en blanco no gasta una llamada de OCR
    assert blank.full_text == ""


def _xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Gastos"
    sheet.append(["Área", "Monto"])
    sheet.append(["Ventas", 1250.5])
    other = workbook.create_sheet("Presupuesto")
    other.append(["Ventas", 10000])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_each_sheet_is_a_page() -> None:
    data = _xlsx_bytes()
    assert detect_content_type(data, "gastos.xlsx") == XLSX
    parsed = parse_document(data, XLSX)
    assert parsed.page_count == 2
    assert parsed.pages[0].text.startswith("Hoja: Gastos")
    assert "Ventas | 1250.5" in parsed.pages[0].text
    assert parsed.pages[1].text.startswith("Hoja: Presupuesto")


def test_pptx_each_slide_is_a_page_with_notes() -> None:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Ingresos Q1"
    slide.placeholders[1].text_frame.text = "Total: 250.000 USD"
    slide.notes_slide.notes_text_frame.text = "Mencionar el margen"
    buffer = io.BytesIO()
    presentation.save(buffer)

    parsed = parse_document(buffer.getvalue(), PPTX)
    assert parsed.page_count == 1
    text = parsed.pages[0].text
    assert "Ingresos Q1" in text and "250.000 USD" in text and "Notas: Mencionar" in text


def test_csv_needs_extension_and_detects_separator() -> None:
    data = b"fecha;monto\n2026-01-02;150.00\n"
    assert detect_content_type(data) is None  # sin extensión no se adivina
    assert detect_content_type(data, "movimientos.csv") == CSV
    assert detect_content_type(data, "movimientos.txt") == TXT
    parsed = parse_document(data, CSV)
    assert parsed.full_text == "fecha | monto\n2026-01-02 | 150.00"


def test_binary_disguised_as_csv_is_rejected() -> None:
    assert detect_content_type(b"\x00\x01\x02binario", "datos.csv") is None


def test_empty_office_files_have_no_text() -> None:
    empty_docx = io.BytesIO()
    docx.Document().save(empty_docx)
    empty_xlsx = io.BytesIO()
    openpyxl.Workbook().save(empty_xlsx)
    for data, content_type in [
        (empty_docx.getvalue(), DOCX),
        (empty_xlsx.getvalue(), XLSX),
        (b"", TXT),
    ]:
        assert parse_document(data, content_type).full_text == ""
