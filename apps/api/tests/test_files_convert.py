"""Conversión de formatos y compresión (funciones puras, sin red ni base de datos)."""

import io
import os
import zipfile

import docx
import openpyxl
import pptx
import pymupdf
import pytest
from PIL import Image

from app.services.files.compress import make_zip, reduce_size
from app.services.files.convert import ConversionError, convert, supported_targets
from app.services.ingestion.filetypes import (
    CSV,
    DOCX,
    JPEG,
    PDF,
    PNG,
    PPTX,
    TXT,
    XLSX,
    ZIP,
    detect_content_type,
)
from app.services.ingestion.parsing import parse_document
from tests.conftest import make_pdf, make_scanned_pdf


def _docx() -> bytes:
    document = docx.Document()
    document.add_heading("Política de viajes", level=1)
    document.add_paragraph("Los viajes se aprueban con 10 días de anticipación.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Categoría"
    table.rows[0].cells[1].text = "Límite"
    document.add_paragraph("Texto posterior a la tabla.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx(sheets: int = 1) -> bytes:
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Gastos"
    first.append(["Área", "Monto"])
    first.append(["Ventas", 1200])
    for n in range(1, sheets):
        workbook.create_sheet(f"Hoja{n + 1}").append(["dato", n])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _png(size: tuple[int, int] = (400, 300), mode: str = "RGB") -> bytes:
    buffer = io.BytesIO()
    Image.new(mode, size, "white").save(buffer, format="PNG")
    return buffer.getvalue()


def test_docx_to_pdf_keeps_text_tables_and_order() -> None:
    out = convert(_docx(), DOCX, "politica.docx", "pdf")
    assert out.filename == "politica.pdf"
    assert out.content_type == PDF and detect_content_type(out.data) == PDF
    text = parse_document(out.data, PDF).full_text
    assert text.index("10 días") < text.index("Categoría") < text.index("Texto posterior")


def test_xlsx_to_csv_single_sheet_and_zip_for_many() -> None:
    single = convert(_xlsx(1), XLSX, "gastos.xlsx", "csv")
    assert single.filename == "gastos.csv" and single.content_type == CSV
    assert single.data.decode("utf-8-sig").splitlines() == ["Área,Monto", "Ventas,1200"]

    many = convert(_xlsx(3), XLSX, "gastos.xlsx", "csv")
    assert many.content_type == ZIP
    names = zipfile.ZipFile(io.BytesIO(many.data)).namelist()
    assert names == ["gastos_Gastos.csv", "gastos_Hoja2.csv", "gastos_Hoja3.csv"]


def test_csv_to_xlsx_converts_numbers() -> None:
    out = convert(b"fecha;monto\n2026-01-02;150.5\n", CSV, "movs.csv", "xlsx")
    sheet = openpyxl.load_workbook(io.BytesIO(out.data)).active
    assert sheet is not None
    assert [c.value for c in sheet[2]] == ["2026-01-02", 150.5]


def test_pptx_to_pdf_one_page_per_slide() -> None:
    presentation = pptx.Presentation()
    for title in ["Q1", "Q2"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Ingresos {title}"
    buffer = io.BytesIO()
    presentation.save(buffer)
    out = convert(buffer.getvalue(), PPTX, "resultados.pptx", "pdf")
    parsed = parse_document(out.data, PDF)
    assert parsed.page_count == 2
    assert "Ingresos Q2" in parsed.pages[1].text


def test_pdf_to_images_zips_multiple_pages() -> None:
    pdf = make_pdf(["<p>Página uno con texto</p>", "<p>Página dos con texto</p>"])
    out = convert(pdf, PDF, "contrato.pdf", "png")
    assert out.content_type == ZIP
    assert zipfile.ZipFile(io.BytesIO(out.data)).namelist() == [
        "contrato_p1.png",
        "contrato_p2.png",
    ]


def test_image_conversions_and_text_targets() -> None:
    assert convert(_png(), PNG, "foto.png", "jpg").content_type == JPEG
    assert detect_content_type(convert(_png(), PNG, "foto.png", "pdf").data) == PDF
    assert convert(b"hola\n\nmundo", TXT, "notas.txt", "docx").content_type == DOCX


def test_unsupported_and_scanned_conversions_explain_why() -> None:
    with pytest.raises(ConversionError, match="Formatos posibles"):
        convert(_png(), PNG, "foto.png", "xlsx")
    with pytest.raises(ConversionError, match="escaneo"):
        convert(make_scanned_pdf("Factura escaneada"), PDF, "scan.pdf", "docx")
    with pytest.raises(ConversionError, match="vacía"):
        convert(_empty_xlsx(), XLSX, "vacio.xlsx", "csv")
    assert "png" not in supported_targets(PNG)  # no se "convierte" a su mismo formato


def _empty_xlsx() -> bytes:
    buffer = io.BytesIO()
    openpyxl.Workbook().save(buffer)
    return buffer.getvalue()


def test_zip_renames_duplicates() -> None:
    out = make_zip([("a.pdf", b"1"), ("a.pdf", b"2"), ("b.txt", b"3")], "lote")
    assert out.filename == "lote.zip"
    assert zipfile.ZipFile(io.BytesIO(out.data)).namelist() == ["a.pdf", "a (1).pdf", "b.txt"]


def test_reduce_size_shrinks_scanned_pdf_and_rejects_office() -> None:
    # Una "foto" grande y con ruido, como una imagen de cámara.
    photo = Image.frombytes("RGB", (2400, 1800), os.urandom(2400 * 1800 * 3))
    buffer = io.BytesIO()
    photo.save(buffer, format="PNG")
    reduced, changed = reduce_size(buffer.getvalue(), PNG, "scan.png")
    assert changed and len(reduced.data) < len(buffer.getvalue())
    assert reduced.filename == "scan_comprimido.jpg"
    assert max(Image.open(io.BytesIO(reduced.data)).size) == 1600

    # Un escaneo real: la foto ocupa una página A4 (~290 DPI efectivos).
    scan = pymupdf.open()
    page = scan.new_page()
    page.insert_image(page.rect, stream=buffer.getvalue())
    scan_bytes = scan.tobytes()
    smaller_pdf, pdf_changed = reduce_size(scan_bytes, PDF, "scan.pdf")
    assert pdf_changed and len(smaller_pdf.data) < len(scan_bytes) / 2

    transparent = _png((50, 50), "RGBA")
    out, _ = reduce_size(transparent, PNG, "logo.png")
    assert out.content_type == PNG  # la transparencia se conserva

    with pytest.raises(ConversionError, match="ZIP"):
        reduce_size(_docx(), DOCX, "politica.docx")
