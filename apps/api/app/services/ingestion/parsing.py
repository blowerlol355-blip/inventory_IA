"""Extracción de texto conservando el número de página.

- PDF: texto nativo con PyMuPDF. Una página sin texto (escaneo) se rasteriza y queda marcada
  para OCR, que hace el pipeline con un LLM con visión.
- DOCX: párrafos y tablas en orden (DOCX no tiene páginas fijas: todo es página 1).
- XLSX: cada hoja es una "página" (la cita apunta a la hoja).
- PPTX: cada diapositiva es una "página" (título, textos, tablas y notas).
- CSV / TXT: todo el texto es la página 1.
- Imágenes: una página marcada para OCR.

Este módulo es síncrono y de CPU: llamarlo con asyncio.to_thread.
"""

import csv
import io
import unicodedata
from dataclasses import dataclass

import docx
import openpyxl
import pptx
import pymupdf
from docx.table import Table
from PIL import Image

from app.services.ingestion.filetypes import (
    CSV,
    DOCX,
    IMAGE_TYPES,
    PDF,
    PPTX,
    TXT,
    XLSX,
    decode_text,
)

MIN_TEXT_CHARS = 20
OCR_DPI = 150
# Lado mayor máximo de la imagen enviada a OCR (el LLM la reduciría de todos modos).
OCR_MAX_SIDE = 2000
# Una hoja de cálculo enorme no aporta a la búsqueda más allá de cierto punto.
MAX_SHEET_ROWS = 5000


class ParsingError(Exception):
    pass


@dataclass
class PageText:
    page: int  # 1-indexado (en Excel, la hoja; en PowerPoint, la diapositiva)
    text: str
    # PNG de la página cuando no tiene texto extraíble y necesita OCR.
    image: bytes | None = None


@dataclass
class ParsedDocument:
    pages: list[PageText]
    page_count: int

    @property
    def ocr_pages(self) -> list[PageText]:
        return [p for p in self.pages if p.image is not None]

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)


def normalize_text(text: str) -> str:
    # NFKC convierte ligaduras tipográficas ("ﬁ" -> "fi") que romperían la búsqueda.
    return unicodedata.normalize("NFKC", text).strip()


def _to_png(image: Image.Image) -> bytes:
    image = image.convert("RGB")
    image.thumbnail((OCR_MAX_SIDE, OCR_MAX_SIDE))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    return buffer.getvalue()


def _parse_pdf(data: bytes) -> ParsedDocument:
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ParsingError("PDF dañado o ilegible") from exc
    pages: list[PageText] = []
    with pdf:
        for index in range(1, pdf.page_count + 1):
            page = pdf[index - 1]
            text = normalize_text(page.get_text("text", sort=True))
            if len(text) >= MIN_TEXT_CHARS:
                pages.append(PageText(page=index, text=text))
            elif not page.get_images() and not page.get_drawings():
                # Página en blanco: no hay nada que transcribir (ahorra una llamada de OCR).
                pages.append(PageText(page=index, text=text))
            else:
                pix = page.get_pixmap(dpi=OCR_DPI)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                pages.append(PageText(page=index, text="", image=_to_png(image)))
        return ParsedDocument(pages=pages, page_count=pdf.page_count)


def _parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParsingError("Word dañado o ilegible") from exc
    blocks: list[str] = []
    # Párrafos y tablas en el orden real del documento.
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in item.rows]
            blocks.append("\n".join(rows))
        elif item.text.strip():
            blocks.append(item.text)
    text = normalize_text("\n\n".join(blocks))
    return ParsedDocument(pages=[PageText(page=1, text=text)], page_count=1)


def _cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def _parse_xlsx(data: bytes) -> ParsedDocument:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParsingError("Excel dañado o ilegible") from exc
    pages: list[PageText] = []
    for index, sheet in enumerate(workbook.worksheets, start=1):
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True, max_row=MAX_SHEET_ROWS):
            cells = [_cell(v) for v in row]
            if any(cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        text = normalize_text(f"Hoja: {sheet.title}\n" + "\n".join(rows)) if rows else ""
        pages.append(PageText(page=index, text=text))
    workbook.close()
    return ParsedDocument(pages=pages, page_count=len(pages))


def _parse_pptx(data: bytes) -> ParsedDocument:
    try:
        presentation = pptx.Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ParsingError("PowerPoint dañado o ilegible") from exc
    pages: list[PageText] = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                blocks.append(shape.text_frame.text.strip())
            if shape.has_table:
                rows = [" | ".join(c.text.strip() for c in r.cells) for r in shape.table.rows]
                blocks.append("\n".join(rows))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            blocks.append("Notas: " + slide.notes_slide.notes_text_frame.text.strip())
        text = normalize_text(f"Diapositiva {index}\n" + "\n\n".join(blocks)) if blocks else ""
        pages.append(PageText(page=index, text=text))
    return ParsedDocument(pages=pages, page_count=len(pages))


def read_csv_rows(text: str) -> list[list[str]]:
    """Filas de un CSV detectando el separador (coma, punto y coma, tabulador)."""
    if not text.strip():
        return []
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], ",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in row)]


def _parse_text(data: bytes, content_type: str) -> ParsedDocument:
    text = decode_text(data)
    if text is None:
        raise ParsingError("El archivo no es texto válido")
    if content_type == CSV:
        text = "\n".join(" | ".join(c.strip() for c in row) for row in read_csv_rows(text))
    return ParsedDocument(pages=[PageText(page=1, text=normalize_text(text))], page_count=1)


def _parse_image(data: bytes) -> ParsedDocument:
    try:
        image = Image.open(io.BytesIO(data))
    except Exception as exc:
        raise ParsingError("Imagen dañada o ilegible") from exc
    return ParsedDocument(pages=[PageText(page=1, text="", image=_to_png(image))], page_count=1)


def parse_document(data: bytes, content_type: str) -> ParsedDocument:
    if content_type == PDF:
        return _parse_pdf(data)
    if content_type == DOCX:
        return _parse_docx(data)
    if content_type == XLSX:
        return _parse_xlsx(data)
    if content_type == PPTX:
        return _parse_pptx(data)
    if content_type in (CSV, TXT):
        return _parse_text(data, content_type)
    if content_type in IMAGE_TYPES:
        return _parse_image(data)
    raise ParsingError(f"Tipo no soportado: {content_type}")
